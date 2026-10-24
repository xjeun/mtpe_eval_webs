import io
from typing import Dict, Optional
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADER_FILL = PatternFill("solid", fgColor="EEF2FF")
BOLD = Font(bold=True)

def _autofit(ws, max_width=80):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(10, max_len + 2), max_width)

def build_excel_bytes(
    df: pd.DataFrame,
    metrics: Dict[str, object],
    meta: Dict[str, object],
    id_col: str,
    src_col: str,
    ref_col: Optional[str],
    hyp_col: str
) -> bytes:
    wb = Workbook()

    # ---- Summary 시트 ----
    ws0 = wb.active
    ws0.title = "Summary"
    ws0["A1"] = "Run Info"
    ws0["A1"].font = Font(bold=True, size=14)

    row = 3
    for k, v in meta.items():
        ws0[f"A{row}"] = str(k)
        ws0[f"B{row}"] = str(v)
        ws0[f"A{row}"].font = BOLD
        row += 1

    row += 1
    ws0[f"A{row}"] = "Metrics"
    ws0[f"A{row}"].font = Font(bold=True, size=14)
    row += 2

    for k, v in metrics.items():
        if isinstance(v, (list, tuple)):
            continue
        ws0[f"A{row}"] = str(k)
        ws0[f"B{row}"] = str(v)
        ws0[f"A{row}"].font = BOLD
        row += 1

    _autofit(ws0, max_width=60)

    # ---- Results 시트 ----
    ws1 = wb.create_sheet("Results")

    cols = [id_col, src_col]
    if ref_col is not None and ref_col in df.columns:
        cols.append(ref_col)
    cols.append(hyp_col)

    # ✅ sent_cometkiwi 포함
    extra_cols = []
    for c in ["sent_bleu", "sent_chrf", "sent_comet", "sent_cometkiwi",
              "prompt_tokens", "completion_tokens", "total_tokens"]:
        if c in df.columns:
            extra_cols.append(c)
    cols += extra_cols

    out_df = df[cols].copy()

    for r_idx, row_vals in enumerate(dataframe_to_rows(out_df, index=False, header=True), start=1):
        ws1.append(row_vals)
        if r_idx == 1:
            for c_idx in range(1, len(cols) + 1):
                cell = ws1.cell(row=1, column=c_idx)
                cell.font = BOLD
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws1.freeze_panes = "A2"

    end_col_letter = ws1.cell(row=1, column=len(cols)).column_letter
    table_ref = f"A1:{end_col_letter}{len(out_df)+1}"

    table = Table(displayName="TranslationTable", ref=table_ref)
    style = TableStyleInfo(
        name="TableStyleLight1",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws1.add_table(table)

    _autofit(ws1, max_width=80)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()